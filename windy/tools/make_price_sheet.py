#!/usr/bin/env python3
"""値段の計算表 (windy/pricing/WINDY_利益計算.xlsx) を作るスクリプト。"""
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

YELLOW = PatternFill("solid", fgColor="FFF2B3")
GRAY = PatternFill("solid", fgColor="EEEEEE")
GREEN = PatternFill("solid", fgColor="D9F2D9")
HEAD = PatternFill("solid", fgColor="333333")
thin = Side(style="thin", color="BBBBBB")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="center")

wb = Workbook()
ws = wb.active
ws.title = "利益計算"
ws.column_dimensions["A"].width = 34
ws.column_dimensions["B"].width = 14
ws.column_dimensions["C"].width = 14
ws.column_dimensions["D"].width = 60


def row(r, label, b=None, c=None, note="", fill=None, fmt=None, bold=False):
    ws.cell(r, 1, label).alignment = WRAP
    for col, v in ((2, b), (3, c)):
        cell = ws.cell(r, col, v)
        cell.border = BOX
        if fill:
            cell.fill = fill
        if fmt:
            cell.number_format = fmt
        if bold:
            cell.font = Font(bold=True)
    ws.cell(r, 4, note).alignment = WRAP
    if bold:
        ws.cell(r, 1).font = Font(bold=True)


ws["A1"] = "WINDY 1枚あたりの利益計算"
ws["A1"].font = Font(size=14, bold=True)
ws["A2"] = "黄色のマスだけ入力してください。緑は自動計算。灰色は公式ページで確認した手数料（変わったら書き換える）。"
ws.merge_cells("A2:D2")

r = 4
for col, t in enumerate(["項目", "半袖T", "ロンT", "説明"], 1):
    c = ws.cell(r, col, t)
    c.fill, c.font = HEAD, Font(color="FFFFFF", bold=True)

YEN = '#,##0"円"'
USD = '"$"#,##0.00'
PCT = "0.0%"
# --- 入力 ---
row(5, "【入力】販売価格（税込）", None, None, "BASEで表示する値段", YELLOW, YEN)
row(6, "【入力】お客さんから受け取る送料", None, None, "送料無料にするなら 0 を入れる", YELLOW, YEN)
row(7, "【入力】Printful 商品代（ボディ＋印刷）", None, None,
    "Printfulの商品ページ・注文画面に出る金額。ドル表示ならドルのまま入れる", YELLOW, "#,##0.00")
row(8, "【入力】Printful 送料（日本向け1枚目）", None, None,
    "注文画面で確認した金額。2枚以上まとめて買われると1枚あたりは下がる", YELLOW, "#,##0.00")
row(9, "【入力】Printful 請求に上乗せされる税など（%）", None, None,
    "未確認。消費税などが上乗せされるか、最初の注文の請求書で確認して入れる（無ければ0）", YELLOW, PCT)
row(10, "【入力】為替レート（1ドル＝何円）", None, None,
    "上の7・8行がドルなら入れる。円で入れた場合は 1 を入れる", YELLOW, "0.00")

# --- 設定（公式確認済み） ---
row(12, "BASE 決済手数料（率）", 0.036, 0.036, "BASE公式 料金ページ（スタンダードプラン）で確認 2026-10-23", GRAY, PCT)
row(13, "BASE 決済手数料（1注文ごとの固定額）", 40, 40, "同上", GRAY, YEN)
row(14, "BASE サービス利用料（率）", 0.03, 0.03, "同上", GRAY, PCT)
row(15, "振込1回で引かれる手数料の合計", 750, 750,
    "BASE公式ヘルプ: 振込手数料250円＋事務手数料500円(振込額2万円未満のとき。2万円以上は0円→250を入れる)", GRAY, YEN)
row(16, "【入力】1回の振込で何枚分まとめて受け取るか", None, None,
    "例: 月1回振込で10枚売れたら 10。空欄なら振込手数料は計算に入れない", YELLOW, "0")

# --- 計算 ---
inputs = lambda c: f"{c}5,{c}7,{c}8,{c}10"
for col in "BC":
    ok = f"COUNT({inputs(col)})=4"
    f = {
        18: f"=IF({ok},{col}5+N({col}6),\"未入力\")",
        19: f"=IF({ok},ROUND({col}18*{col}12+{col}13,0),\"未入力\")",
        20: f"=IF({ok},ROUND({col}18*{col}14,0),\"未入力\")",
        21: f"=IF({ok},ROUND(({col}7+{col}8)*(1+N({col}9))*{col}10,0),\"未入力\")",
        22: f"=IF({ok},IF(N({col}16)>0,ROUND({col}15/{col}16,0),0),\"未入力\")",
        23: f"=IF({ok},{col}18-{col}19-{col}20-{col}21-{col}22,\"未入力\")",
        24: f"=IF({ok},{col}23/{col}18,\"未入力\")",
        25: f"=IF({ok},ROUNDUP(({col}21+{col}22+{col}13)/(1-{col}12-{col}14)-N({col}6),-1),\"未入力\")",
    }
    for rr, formula in f.items():
        ws[f"{col}{rr}"] = formula

row(18, "お客さんの支払い合計", note="販売価格＋送料。BASEの手数料はこの合計にかかる（確信度: 中）", fill=GREEN, fmt=YEN)
row(19, "− BASE 決済手数料", note="合計×3.6%＋40円", fill=GREEN, fmt=YEN)
row(20, "− BASE サービス利用料", note="合計×3%", fill=GREEN, fmt=YEN)
row(21, "− Printful への支払い（円）", note="(商品代＋送料)×(1＋税)×為替", fill=GREEN, fmt=YEN)
row(22, "− 振込手数料（1枚あたり）", note="振込1回の手数料÷まとめる枚数", fill=GREEN, fmt=YEN)
row(23, "＝ 1枚あたりの利益", note="ここがマイナスなら赤字", fill=GREEN, fmt='#,##0"円";[Red]-#,##0"円"', bold=True)
row(24, "利益率", note="利益÷お客さんの支払い合計", fill=GREEN, fmt=PCT)
row(25, "赤字にならない最低の販売価格（目安）", note="これより安くすると1枚売るごとに損をする（10円単位で切り上げ）", fill=GREEN, fmt=YEN)

ws["A27"] = "注意"
ws["A27"].font = Font(bold=True)
notes = [
    "・数字を勝手に入れていません。Printfulの金額はアカウント登録後、商品を選んだ画面で確認して入れてください。",
    "・この表に入っていない費用: サンプル代、返品・不良の対応、広告費、所得税など。",
    "・BASEの手数料は スタンダードプラン（月額0円）の場合。グロースプラン(月額有料)だと率が変わる。",
    "・Printfulの日本向け送料は公式の配送ページに「Tシャツ/トップス・長袖 1点目 $4.95、2点目以降 +$2.00」と表示されていた（2026-10-23、確信度: 中）。必ず注文画面の金額で入れること。",
]
for i, t in enumerate(notes):
    c = ws.cell(28 + i, 1, t)
    ws.merge_cells(start_row=28 + i, start_column=1, end_row=28 + i, end_column=4)
    c.alignment = WRAP
    ws.row_dimensions[28 + i].height = 32

for rr in range(5, 26):
    ws.row_dimensions[rr].height = 30
ws.freeze_panes = "B5"

# --- 値段ごとの比較 ---
ws2 = wb.create_sheet("値段ごとの比較")
ws2["A1"] = "販売価格を変えたときの1枚あたりの利益（他の条件は「利益計算」シートの入力を使う）"
ws2["A1"].font = Font(bold=True)
for col, t in enumerate(["販売価格（税込）", "半袖Tの利益", "ロンTの利益"], 1):
    c = ws2.cell(3, col, t)
    c.fill, c.font = HEAD, Font(color="FFFFFF", bold=True)
    ws2.column_dimensions["ABC"[col - 1]].width = 20
for i, price in enumerate(range(2000, 8001, 500)):
    rr = 4 + i
    ws2.cell(rr, 1, price).number_format = YEN
    for col, src in ((2, "B"), (3, "C")):
        S = "利益計算!"
        ok = f"COUNT({S}{src}7,{S}{src}8,{S}{src}10)=3"
        total = f"($A{rr}+N({S}{src}6))"
        formula = (f"=IF({ok},{total}-ROUND({total}*{S}{src}12+{S}{src}13,0)-ROUND({total}*{S}{src}14,0)"
                   f"-ROUND(({S}{src}7+{S}{src}8)*(1+N({S}{src}9))*{S}{src}10,0)"
                   f"-IF(N({S}{src}16)>0,ROUND({S}{src}15/{S}{src}16,0),0),\"未入力\")")
        c = ws2.cell(rr, col, formula)
        c.number_format = '#,##0"円";[Red]-#,##0"円"'
        c.fill = GREEN
ws2["A18"] = "A列の販売価格は自由に書き換えられます。"

# --- 出典 ---
ws3 = wb.create_sheet("出典")
ws3.column_dimensions["A"].width = 40
ws3.column_dimensions["B"].width = 70
src = [
    ("項目", "出典（2026-10-23 確認）"),
    ("BASE 手数料 3.6%＋40円、サービス利用料3%", "https://thebase.com/price/"),
    ("BASE 振込手数料250円・事務手数料500円(2万円未満)", "https://help.thebase.in/hc/ja/articles/206341302"),
    ("Printful×BASE 連携 手数料表記", "https://www.printful.com/jp/integrations/base"),
    ("Printful 日本向け送料", "https://www.printful.com/jp/shipping"),
]
for i, (a, b) in enumerate(src, 1):
    ws3.cell(i, 1, a)
    ws3.cell(i, 2, b)
ws3["A1"].font = ws3["B1"].font = Font(bold=True)

out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "pricing", "WINDY_利益計算.xlsx")
wb.save(os.path.normpath(out))
print("保存:", os.path.normpath(out))
